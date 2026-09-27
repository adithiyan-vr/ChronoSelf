from pptx import Presentation

def create_pitch_deck():
    prs = Presentation()
    
    # Slide 1: Title & Concept
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "ChronoSelf: Human Potential Optimized"
    slide.placeholders[1].text = "Algorithmic gap arbitration between real-time cognitive behavior and your Ideal Executive Persona."

    # Slide 2: The Problem
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "The Problem"
    tf = slide.placeholders[1].text_frame
    tf.text = "Modern users lose significant productive efficiency to specific friction points and velocity leaks."
    tf.add_paragraph().text = "Key wasters include doomscrolling socials, unfocused sprints, excessive context switching, and reactive email triage."
    tf.add_paragraph().text = "This results in an imbalance between intended cognitive effort and actual time spent on high-leverage tasks."

    # Slide 3: The Solution
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "The Solution"
    tf = slide.placeholders[1].text_frame
    tf.text = "ChronoSelf acts as a centralized Daily Command Hub."
    tf.add_paragraph().text = "It utilizes Bio-Circadian Telemetry to map daily habits against specific target sleep durations and chronotypes (Early Bird or Night Owl)."
    tf.add_paragraph().text = "The system sets optimal peak focus windows to ensure 94% Ultradian Cycle Alignment."

    # Slide 4: Smart Task Replacement Engine
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Smart Task Replacement Engine (Judges Highlight)"
    tf = slide.placeholders[1].text_frame
    tf.text = "The Neural Swap Engine detects planned tasks that present low-value waste or high overhead."
    tf.add_paragraph().text = "It automatically constructs instant cognitive upgrades, replacing tasks like \"Scroll Slack & Social Media\" with deep focus sprints."
    tf.add_paragraph().text = "Accepting these AI heuristic swaps boosts the user's daily Alignment Score and awards ChronoXP."

    # Slide 5: Target Personas & Archetypes
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Target Personas & Archetypes"
    tf = slide.placeholders[1].text_frame
    tf.text = "Student (Academic Polymath): Focuses on exams, research sprints, and intensive study blocks."
    tf.add_paragraph().text = "Working Professional (Strategic High-Performer): Prioritizes deep work, sprint syncs, and cognitive leverage."
    tf.add_paragraph().text = "Executive Staff (Vanguard Visionary): Tracks strategic vision, stakeholder communication, and capital governance."

    # Slide 6: The Gap Breakdown Dashboard
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "The Gap Breakdown Dashboard"
    tf = slide.placeholders[1].text_frame
    tf.text = "Users track their progress against baseline habits using the Gap Breakdown UI."
    tf.add_paragraph().text = "The system monitors Deep Cognitive Work against elapsed low-value friction and drift."
    tf.add_paragraph().text = "It tracks Biometric Energy Resonance to ensure users are operating in their peak cortisol and alpha phases."

    # Slide 7: Technical Implementation
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Technical Implementation"
    tf = slide.placeholders[1].text_frame
    tf.text = "Core Logic: Powered by modular JavaScript (\"app_2.js\") handling state management, local storage telemetry, and dynamic preset loads."
    tf.add_paragraph().text = "Interface: Tailored, responsive web interfaces (\"index_2.html\" and \"onboarding_2.html\") styled with modern Tailwind CSS."
    tf.add_paragraph().text = "Routing: Python-based patch scripts (\"patch_2.py\" and \"patch_index_js_2.py\") strictly manage inter-page routing and telemetry injection."

    prs.save('ChronoSelf_Pitch.pptx')
    print("Presentation 'ChronoSelf_Pitch.pptx' generated successfully!")

if __name__ == '__main__':
    create_pitch_deck()